"""Регистр МРТ-измерений: таблица, которую заполняет нейрорадиолог.

Столбцы приходят из `регистр.csv`, который пишет `Hydrocephalus.ReaderExport`, а
тот берёт их из того же перечня, что и задание специалисту. Здесь столбцы не
объявляются **намеренно**: третий список разъехался бы с первыми двумя, и
разъехался бы молча — заполненный не тот столбец ничем себя не выдаёт.

Зачем вообще книга, если csv уже есть. Специалист заполняет таблицу руками, и
цена опечатки здесь выше цены кода: балл вне шкалы, ноль вместо пропуска или
индекс, посчитанный в уме с ошибкой, портят эталон, по которому потом мерится
всё остальное. Поэтому шкалы закрыты проверкой ввода, индекс считается формулой,
а заполненные заранее ячейки окрашены иначе, чем те, что ждут ввода.

Модулю нужен `openpyxl` — он в необязательной группе зависимостей `register`:
обучение и оценка без него обходятся.
"""

from __future__ import annotations

import csv
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

__all__ = [
    "RegisterError",
    "RegisterSummary",
    "build_register",
]

#: Столбцы, заполненные выгрузкой: специалист их не трогает.
PREFILLED = ("код исследования", "возраст, полных лет", "пол")

#: Столбец, который считается формулой, а не вводится.
COMPUTED = "индекс Эванса (считается)"

#: Два отрезка, из которых считается индекс. Без них книга бессмысленна.
WIDTH_COLUMN = "ширина передних рогов, мм"
SKULL_COLUMN = "поперечник черепа, мм"

#: Закрытые шкалы: столбец — допустимые значения. У миллиметров и углов шкалы
#: нет, и проверка там только помешала бы.
SCALES = {
    "тип чтения (1 первое; 2 повторное)": ("1", "2"),
    "оценка качества (1 пригодно; 2 ограниченно; 3 непригодно)": ("1", "2", "3"),
    "сильвиевы щели, балл": ("0", "1", "2"),
    "сужение верхних борозд, балл": ("0", "1", "2"),
    "фокально расширенные борозды, балл": ("0", "1", "2"),
    "DESH, итоговая оценка": ("0", "1", "2", "3", "4"),
    "Фазекас перивентрикулярно": ("0", "1", "2", "3"),
    "Фазекас глубокое белое вещество": ("0", "1", "2", "3"),
    "перивентрикулярные изменения, балл": ("0", "1", "2"),
    "MTA справа": ("0", "1", "2", "3", "4"),
    "MTA слева": ("0", "1", "2", "3", "4"),
}

#: Коды причин пропуска: те же, что в задании.
REASONS = (
    "1 — нужной серии нет",
    "2 — артефакты не дают измерить",
    "3 — неполный охват",
    "4 — ориентиры не определяются",
    "5 — сопутствующая патология искажает анатомию",
    "9 — иное (напишите в пояснении)",
)

_HOW = (
    ("Регистр МРТ-измерений", True),
    ("", False),
    ("Одна строка — одно чтение одного исследования одним специалистом.", False),
    ("Повторное чтение той же серии — новая строка с типом чтения 2.", False),
    ("", False),
    ("Жёлтые ячейки заполняете вы. Серые уже заполнены или считаются сами.", False),
    ("В ячейку только число; единицы указаны в названии столбца.", False),
    ("«0» значит «признака нет». Пустая ячейка значит «не оценено».", False),
    ("Это разные вещи, и ноль вместо пропуска ставить нельзя.", False),
    ("У каждой пустой ячейки — строка на листе «Контроль качества».", False),
    ("Сомнительный результат записывается как есть и отмечается столбцом", False),
    ("«сомнительный результат» = 1.", False),
    ("", False),
    ("Возраст и пол внесены заранее: из файлов эти теги удалены при", False),
    ("обезличивании, и домысливать их не нужно.", False),
    ("Код серии обязателен: без него измерение не сравнить с автоматическим.", False),
    ("Опись серий каждого исследования — в файле серии.csv рядом со снимками.", False),
    ("", False),
    ("Определения показателей, схемы измерений и порядок работы — в задании.", False),
    ("ФИО, дату рождения и номер истории болезни в файл не вносите.", False),
)

_YELLOW = PatternFill("solid", fgColor="FFF2CC")
_GREY = PatternFill("solid", fgColor="D9D9D9")
_HEAD = PatternFill("solid", fgColor="BDD7EE")
_THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

#: Сколько строк отводится под пропуски на листе контроля качества.
_QUALITY_ROWS = 200


class RegisterError(Exception):
    """Регистр собрать нельзя: в исходной таблице нет нужного столбца."""


@dataclass(frozen=True)
class RegisterSummary:
    """Что получилось: для печати в консоль и для проверки в тестах."""

    rows: int
    columns: int
    editable: int
    path: Path


def build_register(source: Path, workbook: Path) -> RegisterSummary:
    """Собирает книгу регистра из `регистр.csv` выгрузки.

    :param source: путь к `регистр.csv`, записанному выгрузкой.
    :param workbook: куда записать `.xlsx`.
    :returns: сводка по собранной книге.
    :raises RegisterError: если в таблице нет столбцов двух отрезков индекса.
    """
    with source.open(encoding="utf-8-sig", newline="") as handle:
        table = list(csv.reader(handle, delimiter=";"))

    if not table:
        raise RegisterError("Таблица пуста: нет даже строки заголовка.")

    header, data = table[0], table[1:]

    for required in (WIDTH_COLUMN, SKULL_COLUMN):
        if required not in header:
            raise RegisterError(f"В таблице нет столбца «{required}»: индекс считать нечем.")

    book = Workbook()
    sheet = book.active

    if sheet is None:  # pragma: no cover - новая книга всегда с листом
        raise RegisterError("openpyxl не дал листа в новой книге.")

    sheet.title = "Регистр"

    columns = [*header, COMPUTED]
    width_letter = get_column_letter(header.index(WIDTH_COLUMN) + 1)
    skull_letter = get_column_letter(header.index(SKULL_COLUMN) + 1)

    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.fill = _HEAD
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[get_column_letter(index)].width = max(10, min(26, len(name) + 2))

    sheet.row_dimensions[1].height = 58
    sheet.freeze_panes = "B2"

    for line, values in enumerate(data, start=2):
        for index, name in enumerate(columns, start=1):
            if name == COMPUTED:
                # Деление делает книга, а не человек: расхождение в делении
                # потом не отличить от расхождения в измерении.
                cell = sheet.cell(
                    row=line,
                    column=index,
                    value=(
                        f'=IF(OR({width_letter}{line}="",{skull_letter}{line}=""),"",'
                        f"ROUND({width_letter}{line}/{skull_letter}{line},3))"
                    ),
                )
                cell.fill = _GREY
            else:
                cell = sheet.cell(row=line, column=index, value=values[index - 1] or None)
                cell.fill = _GREY if name in PREFILLED else _YELLOW

            cell.border = _THIN
            cell.font = Font(size=10)

    _restrict(sheet, header, len(data))
    _quality_sheet(book)
    _how_sheet(book)

    workbook.parent.mkdir(parents=True, exist_ok=True)
    book.save(workbook)

    editable = sum(1 for name in columns if name not in PREFILLED and name != COMPUTED)

    return RegisterSummary(rows=len(data), columns=len(columns), editable=editable, path=workbook)


def _restrict(sheet: Worksheet, header: list[str], rows: int) -> None:
    """Закрывает шкалы проверкой ввода там, где набор значений конечен."""
    flags = tuple(name for name in header if name.endswith("(0/1)"))
    limits = {name: values for name, values in SCALES.items() if name in header}
    limits.update(dict.fromkeys(flags, ("0", "1")))

    for name, values in limits.items():
        letter = get_column_letter(header.index(name) + 1)
        allowed = ",".join(values)
        rule = DataValidation(type="list", formula1=f'"{allowed}"', allow_blank=True)
        rule.errorTitle = "Вне шкалы"
        rule.error = f"Допустимо: {allowed}. Пустая ячейка значит «не оценено»."
        sheet.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{rows + 1}")


def _quality_sheet(book: Workbook) -> None:
    """Лист пропусков: пустая ячейка без причины считается потерянной."""
    sheet = book.create_sheet("Контроль качества")
    sheet["A1"] = (
        "У каждой пустой ячейки регистра — строка здесь. "
        "Пустая ячейка без причины считается потерянной."
    )
    sheet["A1"].font = Font(bold=True, size=11)
    sheet.merge_cells("A1:D1")

    titles = ("код исследования", "какой столбец не заполнен", "код причины", "пояснение")
    widths = (22, 38, 14, 50)

    for index, (name, width) in enumerate(zip(titles, widths, strict=True), start=1):
        cell = sheet.cell(row=3, column=index, value=name)
        cell.fill = _HEAD
        cell.font = Font(bold=True, size=10)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for line in range(4, 4 + _QUALITY_ROWS):
        for index in range(1, 5):
            cell = sheet.cell(row=line, column=index)
            cell.fill = _YELLOW
            cell.border = _THIN

    rule = DataValidation(type="list", formula1='"1,2,3,4,5,9"', allow_blank=True)
    rule.errorTitle = "Код причины"
    rule.error = "; ".join(REASONS)
    sheet.add_data_validation(rule)
    rule.add(f"C4:C{3 + _QUALITY_ROWS}")

    sheet["F3"] = "Коды причин"
    sheet["F3"].font = Font(bold=True, size=10)

    for shift, text in enumerate(REASONS):
        sheet.cell(row=4 + shift, column=6, value=text).font = Font(size=10)

    sheet.column_dimensions["F"].width = 46


def _how_sheet(book: Workbook) -> None:
    """Короткая памятка на отдельном листе: задание читают не каждый раз."""
    sheet = book.create_sheet("Как заполнять")

    for index, (text, bold) in enumerate(_HOW, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        cell.font = Font(bold=bold, size=12 if bold else 11)

    sheet.column_dimensions["A"].width = 110


def main(argv: list[str] | None = None) -> int:
    """Собирает книгу рядом с `регистр.csv` выгрузки.

    :param argv: аргументы командной строки; первый — каталог выгрузки.
    :returns: код возврата процесса.
    """
    arguments = list(sys.argv[1:] if argv is None else argv)

    if not arguments:
        print("Использование: python -m hydrocephalus_ml.reader_register <каталог выгрузки>")

        return 2

    export = Path(arguments[0])
    summary = build_register(export / "регистр.csv", export / "Регистр МРТ-измерений.xlsx")

    print(f"записан: {summary.path}")
    print(f"строк: {summary.rows}; столбцов: {summary.columns}; заполняемых: {summary.editable}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
