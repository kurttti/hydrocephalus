"""Проверка книги регистра, которую заполняет нейрорадиолог.

Проверяется не вид книги, а те её свойства, из-за которых она вообще собирается:
индекс считается формулой, а не человеком; шкалы закрыты проверкой ввода;
заполненные заранее ячейки отличимы от ждущих ввода; пропуск остаётся пропуском,
а не превращается в ноль.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from hydrocephalus_ml.reader_register import (
    COMPUTED,
    SKULL_COLUMN,
    WIDTH_COLUMN,
    RegisterError,
    build_register,
)

HEADER = (
    "код исследования",
    "код специалиста",
    "тип чтения (1 первое; 2 повторное)",
    "объёмный раздел применим (0/1)",
    "код серии: аксиальная",
    WIDTH_COLUMN,
    SKULL_COLUMN,
    "MTA справа",
    "пустое турецкое седло (0/1)",
    "комментарий",
)

_YELLOW = "00FFF2CC"
_GREY = "00D9D9D9"
_VOLUMETRIC = "объёмный раздел применим (0/1)"


def write_csv(path: Path, header: tuple[str, ...] = HEADER, rows: int = 2) -> Path:
    lines = [";".join(f'"{name}"' for name in header)]

    for number in range(1, rows + 1):
        values = [""] * len(header)
        values[0] = f"S-{number:04d}"

        if _VOLUMETRIC in header:
            values[header.index(_VOLUMETRIC)] = "1"

        lines.append(";".join(values))

    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")

    return path


def test_index_is_a_formula_not_a_number(tmp_path: Path) -> None:
    source = write_csv(tmp_path / "регистр.csv")

    summary = build_register(source, tmp_path / "книга.xlsx")

    sheet = load_workbook(summary.path)["Регистр"]
    assert sheet.cell(row=1, column=sheet.max_column).value == COMPUTED
    formula = sheet.cell(row=2, column=sheet.max_column).value
    assert isinstance(formula, str)
    assert formula.startswith("=IF(")
    assert "ROUND(F2/G2,3)" in formula


def test_prefilled_cells_are_shaded_apart_from_input(tmp_path: Path) -> None:
    # Серое уже заполнено или считается само, жёлтое ждёт ввода. Перепутать их
    # значит либо переписать заданное, либо не заполнить нужное.
    source = write_csv(tmp_path / "регистр.csv")

    summary = build_register(source, tmp_path / "книга.xlsx")

    sheet = load_workbook(summary.path)["Регистр"]
    assert sheet.cell(row=2, column=1).value == "S-0001"
    assert sheet.cell(row=2, column=1).fill.start_color.rgb == _GREY
    assert sheet.cell(row=2, column=4).value == "1"
    assert sheet.cell(row=2, column=4).fill.start_color.rgb == _GREY
    assert sheet.cell(row=2, column=6).fill.start_color.rgb == _YELLOW


def test_age_and_sex_are_not_handed_to_the_reader(tmp_path: Path) -> None:
    # Медиана возраста 74 против 59: возраст в строке выдаёт группу до того, как
    # открыт снимок. В книге его нет, даже если он есть в исходной таблице.
    header = (*HEADER, "возраст, полных лет", "пол")
    source = write_csv(tmp_path / "регистр.csv", header=header)

    summary = build_register(source, tmp_path / "книга.xlsx")

    sheet = load_workbook(summary.path)["Регистр"]
    names = [sheet.cell(row=1, column=index).value for index in range(1, sheet.max_column + 1)]
    assert "возраст, полных лет" not in names
    assert "пол" not in names


def test_an_empty_cell_stays_empty(tmp_path: Path) -> None:
    # Ноль вместо пропуска означал бы «признака нет» там, где признак не
    # оценивали: эти две вещи в анализе нельзя смешивать.
    source = write_csv(tmp_path / "регистр.csv")

    summary = build_register(source, tmp_path / "книга.xlsx")

    sheet = load_workbook(summary.path)["Регистр"]
    assert sheet.cell(row=2, column=6).value is None
    assert sheet.cell(row=2, column=8).value is None


def test_closed_scales_are_restricted(tmp_path: Path) -> None:
    source = write_csv(tmp_path / "регистр.csv")

    summary = build_register(source, tmp_path / "книга.xlsx")

    sheet = load_workbook(summary.path)["Регистр"]
    ranges = {
        str(next(iter(rule.sqref.ranges))): rule.formula1
        for rule in sheet.data_validations.dataValidation
    }
    assert ranges["C2:C3"] == '"1,2"'
    assert ranges["H2:H3"] == '"0,1,2,3,4"'
    assert ranges["I2:I3"] == '"0,1"'


def test_three_sheets_including_the_reasons(tmp_path: Path) -> None:
    source = write_csv(tmp_path / "регистр.csv")

    summary = build_register(source, tmp_path / "книга.xlsx")

    book = load_workbook(summary.path)
    assert book.sheetnames == ["Регистр", "Контроль качества", "Как заполнять"]
    assert book["Контроль качества"]["F4"].value is not None


def test_summary_counts_editable_columns(tmp_path: Path) -> None:
    source = write_csv(tmp_path / "регистр.csv", rows=5)

    summary = build_register(source, tmp_path / "книга.xlsx")

    assert summary.rows == 5
    assert summary.columns == len(HEADER) + 1
    # Не заполняются два столбца — код и отметка объёма, — и ещё считаемый.
    assert summary.editable == len(HEADER) - 2


def test_a_table_without_the_two_segments_is_refused(tmp_path: Path) -> None:
    header = tuple(name for name in HEADER if name != WIDTH_COLUMN)
    source = write_csv(tmp_path / "регистр.csv", header=header)

    with pytest.raises(RegisterError, match="ширина передних рогов"):
        build_register(source, tmp_path / "книга.xlsx")


def test_an_empty_table_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "регистр.csv"
    source.write_text("", encoding="utf-8-sig")

    with pytest.raises(RegisterError, match="пуста"):
        build_register(source, tmp_path / "книга.xlsx")
